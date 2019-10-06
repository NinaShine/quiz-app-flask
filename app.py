from flask import Flask,request, render_template, flash, redirect, url_for,session, logging
from flask_mysqldb import MySQL 
from wtforms import Form, StringField, TextAreaField, PasswordField, validators, DateTimeField, BooleanField, IntegerField
from flask_wtf import FlaskForm
from flask_wtf.file import FileField, FileRequired, FileAllowed
from passlib.hash import sha256_crypt
from functools import wraps
from werkzeug.utils import secure_filename
from docx import Document
from coolname import generate_slug
from datetime import timedelta, datetime
from flask_mail import Mail, Message
from threading import Thread
from flask import render_template_string
from itsdangerous import URLSafeTimedSerializer
from validate_email import validate_email
import random
import json

app = Flask(__name__)
app.secret_key= 'huihui'

#Config MySQL
app.config['MYSQL_HOST'] = 'localhost'
app.config['MYSQL_USER'] = 'nubaf'
app.config['MYSQL_PASSWORD'] = 'nubafgg'
app.config['MYSQL_DB'] = 'flask'
app.config['MYSQL_CURSORCLASS'] = 'DictCursor'


app.config.update(
	DEBUG=True,
	MAIL_SERVER='smtp.gmail.com',
	MAIL_PORT=465,
	MAIL_USE_SSL=True,
	MAIL_USERNAME = 'nickqwerty76@gmail.com',
	MAIL_PASSWORD = 'Zeitgeist77'
	)
mail = Mail(app)

def asynch(f):
	@wraps(f)
	def wrapper(*args, **kwargs):
		thr = Thread(target=f, args=args, kwargs=kwargs)
		thr.start()
	return wrapper

@asynch
def send_async_email(app, msg):
	with app.app_context():
		mail.send(msg)


htmlbody='''
Your account on <b>The Best</b> Quiz App was successfully created.
Please click the link below to confirm your email address and
activate your account:
  
<a href="{{ confirm_url }}">{{ confirm_url }}</a>
 
--
Questions? Comments? Email nickqwerty76@gmail.com.
'''

@app.route('/sendmail', methods = ['GET','POST'])
def send_email(recipients,html_body):
	try:
		msg = Message('Confirm Your Email Address',
		  sender="nickqwerty76@gmail.com",
		  recipients=recipients)
		# msg.body = "Yo!\nHave you heard the good word of Python???"
		msg.html = html_body
		send_async_email(app, msg)
		# return 'Mail sent!'
		return
	except Exception as e:
		# return(str(e))
		return


def send_confirmation_email(user_email):
	confirm_serializer = URLSafeTimedSerializer(app.config['SECRET_KEY'])
 
	confirm_url = url_for('confirm_email',
		token=confirm_serializer.dumps(user_email, salt='email-confirmation-salt'),
		_external=True)
 
	html = render_template_string(htmlbody, confirm_url=confirm_url)

	send_email([user_email], html)


#init Mysql
mysql = MySQL(app)

@app.before_request
def make_session_permanent():
	session.permanent = True
	app.permanent_session_lifetime = timedelta(minutes=5)


def is_logged(f):
	@wraps(f)
	def wrap(*args, **kwargs):
		if 'logged_in' in session:
			return f(*args, **kwargs)
		else:
			flash('Unauthorized, Please login','danger')
			return redirect(url_for('login'))
	return wrap


def doctodict(filepath):
	document = Document(filepath)
	data={}
	count=1
	for table in document.tables:
		temp = {}
		for rowNo,_ in enumerate(table.rows):
			temp[table.cell(rowNo, 0).text]=table.cell(rowNo, 1).text
		data[count] = temp
		count+=1
 
	return data


class RegisterForm(Form):
	name = StringField('Name', [validators.Length(min=3, max=50)])
	username = StringField('Username', [validators.Length(min=4,max=25)])
	email = StringField('Email', [validators.Length(min=6,max=50)])
	password = PasswordField('Password', [
			validators.DataRequired(),
			validators.EqualTo('confirm', message="Password do not match")
		])
	confirm = PasswordField('Confirm Password')


class UploadForm(FlaskForm):
	doc = FileField('Docx Upload', validators=[FileRequired()])
	start_date_time = DateTimeField('Start Date & Time')
	end_date_time = DateTimeField('End Date & Time')
	show_result = BooleanField('Show Result after completion')
	duration = IntegerField('Duration')
	password = StringField('Test Password', [validators.Length(min=3, max=6)])


class TestForm(Form):
	test_id = StringField('Test ID')
	password = PasswordField('Test Password')


@app.route('/')
def index():
	return render_template('layout.html')

@app.route('/register', methods=['GET','POST'])
def register():
	form = RegisterForm(request.form)
	if request.method == 'POST' and form.validate():
		name = form.name.data 
		email = form.email.data

		# check if email is valid, verify

		# is_valid = validate_email(email,verify=True)
		# if is_valid == False:
		# 	flash('Wrong email','danger')
		# do something

		username = form.username.data
		password = sha256_crypt.encrypt(str(form.password.data))
		cur = mysql.connection.cursor()
		cur.execute('INSERT INTO users(username,name,email, password,confirmed) values(%s,%s,%s,%s,0)', (username,name, email, password))
		mysql.connection.commit()
		cur.close()
		send_confirmation_email(email)
		flash('Thanks for registering!  Please check your email to confirm your email address.', 'success')
		return redirect(url_for('index')) 
		# change in login function to redirect to warning page

	return render_template('register.html', form=form)

	
@app.route('/login', methods=['GET','POST'])
def login():
	if request.method == 'POST':
		username = request.form['username']
		password_candidate = request.form['password']
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT * from users where username = %s' , [username])
		if results > 0:
			data = cur.fetchone()
			password = data['password']
			confirmed = data['confirmed']
			if confirmed == 0:
				error = 'Please confirm email before logging in'
				return render_template('login.html', error=error)
			if sha256_crypt.verify(password_candidate, password):
				session['logged_in'] = True
				session['username'] = username
				return redirect(url_for('dashboard'))
			else:
				error = 'Invalid password'
				return render_template('login.html', error=error)
			cur.close()
		else:
			error = 'Username not found'
			return render_template('login.html', error=error)
	return render_template('login.html')


@app.route('/dashboard')
@is_logged
def dashboard():
	return render_template('dashboard.html')


@app.route('/logout')
def logout():
	session.clear()
	flash('Successfully logged out', 'success')
	return redirect(url_for('index'))


@app.route('/create-test', methods = ['GET', 'POST'])
@is_logged
def create_test():
	form = UploadForm()
	if request.method == 'POST' and form.validate_on_submit():
		f = form.doc.data
		filename = secure_filename(f.filename)
		f.save('questions/' + filename)
		cur = mysql.connection.cursor()
		d = doctodict('questions/' + f.filename.replace(' ', '_').replace('(','').replace(')',''))
		test_id = generate_slug(2)
		for no, data in d.items():
			marks = data['((MARKS)) (1/2/3...)']
			a = data['((OPTION_A))']
			b = data['((OPTION_B))']
			c = data['((OPTION_C))']
			d = data['((OPTION_D))']
			question = data['((QUESTION))']
			correct_ans = data['((CORRECT_CHOICE)) (A/B/C/D)']
			explanation = data['((EXPLANATION)) (OPTIONAL)']
			cur.execute('INSERT INTO questions(test_id,qid,q,a,b,c,d,ans,marks) values(%s,%s,%s,%s,%s,%s,%s,%s,%s)', 
				(test_id,no,question,a,b,c,d,correct_ans,marks))
			mysql.connection.commit()
		start_date_time = form.start_date_time.data
		end_date_time = form.end_date_time.data
		show_result = form.show_result.data
		duration = form.duration.data
		password = form.password.data
		cur.execute('INSERT INTO teachers (username, test_id, start, end, duration, show_ans, password) values(%s,%s,%s,%s,%s,%s,%s)',
			(dict(session)['username'], test_id, start_date_time, end_date_time, duration, show_result, password))
		mysql.connection.commit()
		cur.close()
	return render_template('create_test.html' , form = form)


@app.route('/give-test/<testid>', methods=['GET','POST'])
@is_logged
def test(testid):
	global duration,marked_ans
	if request.method == 'GET':
		try:
			data = {'duration': duration, 'marks': '', 'q': '', 'a': "", 'b':"",'c':"",'d':"" }
			return render_template('quiz.html' ,**data, answers=marked_ans)
		except:
			return redirect(url_for('give_test'))
	else:
		cur = mysql.connection.cursor()
		flag = request.form['flag']
		if flag == 'get':
			num = request.form['no']
			results = cur.execute('SELECT * from questions where test_id = %s and qid =%s',(testid, num))
			if results > 0:
				data = cur.fetchone()
				del data['ans']
				cur.close()
				return json.dumps(data)
		elif flag=='mark':
			qid = request.form['qid']
			ans = request.form['ans']
			results = cur.execute('SELECT * from students where test_id =%s and qid = %s and username = %s', (testid, qid, session['username']))
			if results > 0:
				cur.execute('UPDATE students set ans = %s where test_id = %s and qid = %s and username = %s', (testid, qid, session['username']))
			else:
				cur.execute('INSERT INTO students values(%s,%s,%s,%s)', (session['username'], testid, qid, ans))
			mysql.connection.commit()
			cur.close()
		elif flag=='time':
			time_left = request.form['time']
			cur.execute('UPDATE studentTestInfo set time_left=SEC_TO_TIME(%s) where test_id = %s and username = %s', (time_left, testid, session['username']))
			mysql.connection.commit()
			cur.close()
		else:
			cur.execute('UPDATE studentTestInfo set completed=true and time_left=sec_to_time(0) where test_id = %s and username = %s', (testid, session['username']))
			mysql.connection.commit()
			cur.close()


@app.route("/give-test", methods = ['GET', 'POST'])
@is_logged
def give_test():
	global duration, marked_ans	
	form = TestForm(request.form)
	if request.method == 'POST' and form.validate():
		test_id = form.test_id.data
		password_candidate = form.password.data
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT * from teachers where test_id = %s', [test_id])
		if results > 0:
			data = cur.fetchone()
			password = data['password']
			duration = data['duration']
			start = data['start']
			end = data['end']
			if password == password_candidate:
				now = datetime.now()
				now = now.strftime("%Y/%m/%d %H:%M:%S")
				if datetime.strptime(start,"%Y/%m/%d %H:%M:%S") < now and datetime.strptime(end,"%Y/%m/%d %H:%M:%S") > now:
					results = cur.execute('SELECT time_to_sec(time_left) as time_left,completed from studentTestInfo where username = %s and test_id = %s', (session['username'], test_id))
					if results > 0:
						results = cur.fetchone()
						is_completed = results['completed']
						if is_completed == 0:
							time_left = results['time_left']
							if time_left <= duration:
								duration = time_left
								results = cur.execute('SELECT * from students where username = %s and test_id = %s', (session['username'], test_id))
								marked_ans = {}
								if results > 0:
									results = cur.fetchall()
									for row in results:
										marked_ans[row['qid']] = row['ans']
									marked_ans = json.dumps(marked_ans)
									#return redirect(url_for('test' , testid = test_id))
						else:
							flash('Test already given', 'success')
							return redirect(url_for('give_test'))
					else:
						cur.execute('INSERT into studentTestInfo (username, test_id,time_left) values(%s,%s,SEC_TO_TIME(%s))', (session['username'], test_id, duration))
						mysql.connection.commit()
				return redirect(url_for('test' , testid = test_id))
			else:
				flash('Invalid password', 'danger')
				return redirect(url_for('give_test'))
		flash('Invalid testid', 'danger')
		return redirect(url_for('give_test'))
		cur.close()
	return render_template('give_test.html', form = form)


@app.route('/randomize', methods = ['POST'])
def random_gen():
	if request.method == "POST":
		id = request.form['id']
		print(id)
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT count(*) from questions where test_id = %s', [id])
		if results > 0:
			data = cur.fetchone()
			total = data['count(*)']
			nos = list(range(1,int(total)+1))
			random.Random(id).shuffle(nos)
			print(nos)
			cur.close()
			return json.dumps(nos)


@app.route('/<username>/<testid>')
@is_logged
def check_result(username, testid):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT * FROM teachers where test_id = %s', [testid])
		if results>0:
			results = cur.fetchone()
			check = results['show_ans']
			if check == 1:
				results = cur.execute('SELECT q,marks,questions.ans as correct, students.ans as marked,a,b,c,d from students,questions where username = %s and students.test_id = questions.test_id and students.test_id = %s and students.qid=questions.qid', (username, testid))
				if results > 0:
					results = cur.fetchall()
					return render_template('tests_result.html', results= results)
			else:
				flash('You are not authorized to check the result', 'danger')
	else:
		return redirect(url_for('dashboard'))
		
@app.route('/<username>/tests-given')
@is_logged
def tests_given(username):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results = cur.execute('select distinct(test_id) from students where username = %s', [username])
		results = cur.fetchall()
		return render_template('tests_given.html', tests=results)
	else:
		flash('You are not authorized', 'danger')
		return redirect(url_for('dashboard'))


@app.route('/<username>/tests-created')
@is_logged
def tests_created(username):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results = cur.execute('select * from teachers where username = %s', [username])
		results = cur.fetchall()
		return render_template('tests_created.html', tests=results)
	else:
		flash('You are not authorized', 'danger')
		return redirect(url_for('dashboard'))


@app.route('/confirm/<token>')
def confirm_email(token):
	try:
		confirm_serializer = URLSafeTimedSerializer(app.config['SECRET_KEY'])
		email = confirm_serializer.loads(token, salt='email-confirmation-salt', max_age=3600)
	except:
		flash('The confirmation link is invalid or has expired.', 'error')
		return redirect(url_for('login'))

	cur = mysql.connection.cursor()
	results = cur.execute('SELECT * from users where email = %s' , [email])
	if results > 0:
		data = cur.fetchone()
		email_confirmed = data['confirmed']
		if email_confirmed:
			flash('Account already confirmed. Please login.', 'info')
		else:
			results = cur.execute('UPDATE users SET confirmed = 1 where email = %s' , [email])
			mysql.connection.commit()
			cur.close()
			flash('Thank you for confirming your email address!', 'success')
	
		return redirect(url_for('index'))


if __name__ == "__main__":
	app.run(debug=True)
